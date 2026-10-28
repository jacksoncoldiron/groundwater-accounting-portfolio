"""Adapter: the analyst's MLRP co-benefits matrix (MLRP_Benefits_Matrix.xlsx).

The workbook rates how each grantee basin (and the literature) quantifies each
co-benefit on a five-level "measurement state" scale defined by the analyst:

    Not found · Named · Screened / scored · Practice monitored · Outcome quantified

On the Overview sheet the state is carried by cell fill, keyed to a legend row
on the same sheet. This adapter reads the legend fills rather than hard-coding
them, so a re-shaded workbook stays correct. The Detailed sheet supplies the
state sentence and the supporting components for each cell.

No state is assigned here -- every value is the analyst's own rating.
"""

from __future__ import annotations

import re

from common import BuildReport, active_basins, resolve_source

SOURCE_CANDIDATES = ["MLRP_Benefits_Matrix.xlsx"]

COMPONENTS = {
    "Measurement state": "state_text",
    "Data used": "data",
    "Quantification process": "process",
    "Metric / output": "metric",
    "Decision informed & stage": "decision",
    "Gaps, caveats & open questions": "gaps",
    "Sources": "sources",
}


def _txt(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"[ \t]+", " ", str(v)).strip()
    return s or None


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def _split_dpm(cell: str | None) -> dict:
    out = {"data": None, "process": None, "metric": None}
    if not cell:
        return out
    for line in cell.splitlines():
        m = re.match(r"\s*(Data|Process|Metric)\s*:\s*(.*)", line)
        if m:
            out[m.group(1).lower()] = m.group(2).strip() or None
    return out


def ingest(report: BuildReport, basin_registry: dict) -> dict:
    import openpyxl

    path = resolve_source(SOURCE_CANDIDATES)
    report.note(f"MLRP matrix source: {path.name}")
    wb = openpyxl.load_workbook(path, data_only=True)
    ov = wb["Overview Matrix"]

    # Legend: the row whose second cell reads "Measurement state:".
    levels, fill_to_level, header_row = [], {}, None
    for row in ov.iter_rows():
        label = _txt(row[1].value) if len(row) > 1 else None
        if label and label.lower().startswith("measurement state"):
            for i, c in enumerate(row[2:]):
                name = _txt(c.value)
                if not name:
                    continue
                rgb = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
                levels.append({"id": _slug(name), "label": name, "rank": len(levels)})
                if rgb:
                    fill_to_level[rgb] = _slug(name)
        if label == "Benefit":
            header_row = row
            break
    if not levels or header_row is None:
        raise ValueError("MLRP matrix: legend or header row not found")

    by_short = {b["short_name"].lower(): b for b in active_basins(basin_registry)}
    columns = []
    for c in header_row[2:]:
        name = _txt(c.value)
        if not name:
            continue
        if name.lower() == "literature":
            columns.append({"id": "literature", "label": "Literature", "kind": "literature"})
        else:
            b = by_short.get(name.lower())
            if not b:
                report.warn(f"MLRP matrix column {name!r} matches no active basin; skipped")
                columns.append(None)
                continue
            columns.append({"id": b["id"], "label": b["short_name"], "kind": "basin"})

    benefits, group = [], None
    for row in ov.iter_rows(min_row=header_row[0].row + 1):
        g, name = _txt(row[0].value), _txt(row[1].value)
        if g:
            group = g
        if not name or name.startswith("Source:"):
            continue
        cells = {}
        for col, c in zip(columns, row[2:]):
            if not col:
                continue
            rgb = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
            state = fill_to_level.get(rgb)
            if state is None:
                report.warn(f"MLRP matrix: unrecognised fill {rgb} at {c.coordinate}")
            cells[col["id"]] = {"state": state, **_split_dpm(c.value)}
        benefits.append({"id": _slug(name), "name": name, "group": group, "cells": cells})

    # Detailed sheet: same benefits, one row per component.
    det = wb["Detailed Matrix"]
    current = None
    by_name = {b["name"]: b for b in benefits}
    det_cols = None
    for row in det.iter_rows(values_only=True):
        if row[1] == "Benefit" and row[2] == "Component":
            det_cols = [ _txt(v) for v in row[3:] ]
            continue
        if det_cols is None:
            continue
        if _txt(row[1]):
            current = by_name.get(_txt(row[1]))
            if current is None:
                break  # past the benefit blocks (references / row index)
        key = COMPONENTS.get(_txt(row[2]) or "")
        if not current or not key:
            continue
        for label, value in zip(det_cols, row[3:]):
            if not label:
                continue
            cid = "literature" if label.lower() == "literature" else (
                by_short.get(label.lower(), {}).get("id"))
            cell = current["cells"].get(cid)
            if cell is None:
                continue
            text = _txt(value)
            if key == "state_text":
                cell["state_text"] = text
            elif key == "sources":
                cell["sources"] = text
                cell["finding_ids"] = sorted(
                    {int(n) for grp in re.findall(r"IDs?\s+([\d,\s–-]+)", text or "")
                     for n in re.findall(r"\d+", grp)}
                )
            else:
                cell.setdefault("detail", {})[key] = text

    note = None
    for row in ov.iter_rows(values_only=True):
        if row[0] and str(row[0]).startswith("Source:"):
            note = _txt(row[0])
    report.note(f"MLRP matrix: {len(benefits)} benefits x {len([c for c in columns if c])} columns")
    return {
        "levels": levels,
        "columns": [c for c in columns if c],
        "benefits": benefits,
        "source_note": note,
    }
